"""Monika: sprawy, dokumenty z maili, zespół, pisma .docx, odsetki, bramka narzędzi."""

import io
import json
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from kancelaria_agent import guards, odsetki
from kancelaria_agent import project_tools as pt
from kancelaria_agent.agent import Agent
from kancelaria_agent.documents import build_docx, docx_text, read_document
from kancelaria_agent.integrations import IncomingMail, build_mock
from kancelaria_agent.planner import PROJECT_UPDATE, PlannedAction
from kancelaria_agent.planner_agent import decide, parse_reply, system_prompt
from kancelaria_agent.projects import safe_filename

BASIA = "basia@kancelariabil.pl"
ANETA = "aneta@kancelariabil.pl"


class ScriptedPlanner:
    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def plan(self, kind, payload):
        self.calls.append((kind, payload))
        resp = self.responses.pop(0) if self.responses else []
        return resp(payload) if callable(resp) else resp


def answer(to, text="Gotowe", files=None, status="gotowe", klient="Alfa", sprawa="umowa spolki"):
    p = {"to": to, "subject": "Re: x", "body": text}
    if files:
        p["project_attachments"] = files
    return [PlannedAction(PROJECT_UPDATE, {"klient": klient, "projekt": sprawa, "status": status}),
            PlannedAction(guards.SEND_MAIL, p)]


@pytest.fixture
def env(make_env):
    def _make(mails, planner, **ov):
        cfg, store, audit, _, _ = make_env(dry_run=False, **ov)
        integ = build_mock(mails)
        return cfg, store, integ, Agent(cfg, store, audit, integ, planner)
    return _make


# ---------- zespół ----------

def test_reply_goes_to_team_member_who_wrote(env):
    m = IncomingMail("<a@x>", ANETA, "Umowa Alfa", "Przygotuj projekt", trusted=True)
    planner = ScriptedPlanner(lambda p: answer(p["sender"]))
    cfg, store, integ, agent = env([m], planner)
    agent.tick()
    assert [s["to"] for s in integ.mail.sent] == [[ANETA]]
    log = store.project_log(store.db.execute("SELECT id FROM projects").fetchone()[0])
    assert log[0]["role"] == "zespol" and log[0]["text"].startswith(ANETA)


def test_mail_outside_team_blocked_even_if_planner_tries(env):
    m = IncomingMail("<b@x>", BASIA, "Wezwanie", "Napisz do dłużnika", trusted=True)
    planner = ScriptedPlanner(answer("dluznik@firma.pl"))
    cfg, store, integ, agent = env([m], planner, guardrails={"allowed_recipients": ["@firma.pl", BASIA]})
    agent.tick()
    assert integ.mail.sent == []
    events = [json.loads(l) for l in cfg.audit_log.read_text(encoding="utf-8").splitlines()]
    assert any(e["event"] == "mail_blocked" for e in events)


def test_cc_outside_team_blocked(env):
    m = IncomingMail("<c@x>", BASIA, "X", "Y", trusted=True)
    acts = answer(BASIA)
    acts[1].payload["cc"] = ["klient@firma.pl"]
    cfg, store, integ, agent = env([m], ScriptedPlanner(acts))
    agent.tick()
    assert integ.mail.sent == []


def test_empty_team_sends_nothing(env):
    m = IncomingMail("<d@x>", BASIA, "X", "Y", trusted=True)
    cfg, store, integ, agent = env([m], ScriptedPlanner(answer(BASIA)), team=[])
    agent.tick()
    assert integ.mail.sent == []


# ---------- dokumenty z maila ----------

def test_attachments_saved_in_case_and_listed_for_planner(env):
    pdf = b"%PDF-1.4 test"
    m = IncomingMail("<e@x>", BASIA, "Analiza umowy", "Przeanalizuj", trusted=True,
                     attachments=("Umowa spółki – Alfa.pdf",), files=(("Umowa spółki – Alfa.pdf", pdf),))
    planner = ScriptedPlanner(answer(BASIA))
    cfg, store, integ, agent = env([m], planner)
    agent.tick()
    saved = planner.calls[0][1]["saved_files"]
    assert len(saved) == 1 and saved[0].startswith("dokumenty/") and saved[0].endswith("umowa-spolki-alfa.pdf")
    p = store.get_project(planner.calls[0][1]["project_id"])
    assert (agent.projects.dir(p) / saved[0]).read_bytes() == pdf
    assert not list((cfg.projects_dir.parent / "przychodzace").iterdir())   # katalog przejściowy sprzątnięty


def test_gmail_keeps_document_bytes_only_for_team():
    from email.message import EmailMessage

    from kancelaria_agent.integrations.gmail import parse_message

    def mk(sender, domain):
        m = EmailMessage()
        m["From"], m["To"], m["Subject"], m["Message-ID"] = sender, "monika@kancelariabil.pl", "x", "<1@x>"
        m["Authentication-Results"] = f"mx.google.com; dmarc=pass (p=NONE) header.from={domain}"
        m.set_content("x")
        m.add_attachment(b"%PDF-1.4", maintype="application", subtype="pdf", filename="pozew.pdf")
        m.add_attachment(b"MZ", maintype="application", subtype="octet-stream", filename="wirus.exe")
        return bytes(m)
    team = (BASIA, ANETA)
    mine = parse_message(mk(ANETA, "kancelariabil.pl"), team, "f")
    assert mine.trusted and [n for n, _ in mine.files] == ["pozew.pdf"]
    other = parse_message(mk("obcy@evil.com", "evil.com"), team, "f")
    assert not other.trusted and other.files == ()
    spoofed = parse_message(mk(BASIA, "evil.com"), team, "f")   # podszycie się pod adres zespołu
    assert not spoofed.trusted and spoofed.files == ()


def test_docx_attachment_goes_out_and_guard_limits_dirs(env, make_env):
    m = IncomingMail("<f@x>", BASIA, "Umowa", "Projekt umowy", trusted=True)

    def make_file(payload):
        p = agent.store.get_project(payload["project_id"])
        rel = pt.create_docx(agent.projects.dir(p), "Projekt umowy", "# Umowa\n\n§ 1\nStrony …")
        return answer(BASIA, files=[rel])
    planner = ScriptedPlanner(make_file)
    cfg, store, integ, agent = env([m], planner)
    agent.tick()
    att = integ.mail.sent[0]["attachments"]
    assert len(att) == 1 and att[0].endswith(".docx") and "/pisma/" in att[0]
    cfg2, *_ = make_env()
    assert guards.attachments_problem(cfg2, ["x/inne/a.docx"])


# ---------- pisma .docx i odczyt ----------

def test_build_docx_and_read_back(tmp_path):
    target = build_docx("# Wezwanie do zapłaty\n>> Warszawa, 20.10.2026\n\n§ 1\nKwota **1 000 zł**.\n"
                        "- punkt\n| Lp. | Faktura |\n|---|---|\n| 1 | FV/1 |\n---\nStrona 2", tmp_path / "w.docx")
    text = docx_text(target)
    assert "Wezwanie do zapłaty" in text and "§ 1" in text and "1 | FV/1" in text and "Strona 2" in text


def test_docx_from_template_keeps_header_and_replaces_body(tmp_path):
    from docx import Document
    tpl = Document()
    tpl.sections[0].header.paragraphs[0].text = "Kancelaria Adwokacka – nagłówek"
    tpl.add_paragraph("STARA TREŚĆ WZORU")
    tpl.save(tmp_path / "wzor.docx")
    out = build_docx("Nowa treść", tmp_path / "out.docx", tmp_path / "wzor.docx")
    doc = Document(str(out))
    assert doc.sections[0].header.paragraphs[0].text == "Kancelaria Adwokacka – nagłówek"
    body = docx_text(out)
    assert "Nowa treść" in body and "STARA" not in body


def test_create_docx_versions_and_template_bounds(tmp_path):
    case, wzory = tmp_path / "sprawa", tmp_path / "wzory"
    wzory.mkdir()
    build_docx("Wzór", wzory / "umowa.docx")
    a = pt.create_docx(case, "Umowa spółki", "A", wzory, "umowa.docx")
    b = pt.create_docx(case, "Umowa spółki", "B")
    assert a == "pisma/umowa-spolki.docx" and b == "pisma/umowa-spolki-v2.docx"
    for bad in ("../sprawa/pisma/umowa-spolki.docx", "/etc/passwd", "brak.docx"):
        with pytest.raises(pt.ToolError):
            pt.create_docx(case, "x", "y", wzory, bad)


def test_read_file_only_case_dirs(tmp_path):
    (tmp_path / "dokumenty").mkdir()
    (tmp_path / "dokumenty" / "a.txt").write_text("treść", encoding="utf-8")
    assert pt.read_file(tmp_path, "dokumenty/a.txt")[0]["text"] == "treść"
    (tmp_path / "sekret.txt").write_text("x")
    for bad in ("sekret.txt", "../x", "dokumenty/../sekret.txt"):
        with pytest.raises(pt.ToolError):
            pt.read_file(tmp_path, bad)


def test_read_pdf_text(tmp_path):
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(200, 200)
    with open(tmp_path / "pusty.pdf", "wb") as f:
        w.write(f)
    try:
        blocks = read_document(tmp_path / "pusty.pdf")
    except Exception as e:   # bez pdftoppm na maszynie testowej – czytelny błąd, nie wyjątek SDK
        assert "pdftoppm" in str(e)
        return
    assert blocks[0]["type"] == "text" and "skan" in blocks[0]["text"]


def test_safe_filename():
    assert safe_filename("Umowa spółki – Alfa (2).PDF") == "umowa-spolki-alfa-2.pdf"
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("....") == "plik"


def test_templates_import_and_listing(tmp_path, make_env, monkeypatch):
    from kancelaria_agent.__main__ import main
    cfg, *_ = make_env()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Wzory/Spółki/Umowa spółki z o.o..docx", b"PK..")
        z.writestr("Wzory/~$blokada.docx", b"x")
        z.writestr("Wzory/skrypt.sh", b"x")
        z.writestr("../evil.docx", b"x")
    zp = tmp_path / "wzory.zip"
    zp.write_bytes(buf.getvalue())
    assert main(["--config", str(cfg.source_path), "wzory-import", str(zp)]) == 0
    names = [n for n, _ in pt.list_templates(cfg.templates_dir)]
    assert names == ["wzory/spolki/umowa-spolki-z-o-o.docx"]
    assert pt.list_templates(cfg.templates_dir, "umowa spółki") and not pt.list_templates(cfg.templates_dir, "pozew")


# ---------- odsetki ----------

def test_interest_statutory_delay_across_rate_change():
    w = odsetki.oblicz("10000", date(2025, 10, 1), date(2025, 10, 31), "opoznienie")
    # 1–8.10: 4,75+5,5 = 10,25%; 9–31.10: 4,50+5,5 = 10,00%
    assert [(o.stopa, o.dni) for o in w.odcinki] == [(Decimal("10.25"), 8), (Decimal("10.00"), 23)]
    assert w.suma == Decimal("22.47") + Decimal("63.01")
    assert "art. 481 § 2 k.c." in w.opis()


def test_interest_commercial_uses_half_year_rate():
    w = odsetki.oblicz("1000", date(2026, 6, 1), date(2026, 7, 31), "handlowe")
    assert [(o.stopa, o.dni) for o in w.odcinki] == [(Decimal("14.00"), 30), (Decimal("13.75"), 31)]
    w2 = odsetki.oblicz("1000", date(2026, 6, 1), date(2026, 6, 30), "handlowe", podmiot_leczniczy=True)
    assert w2.odcinki[0].stopa == Decimal("12.00")
    old = odsetki.oblicz("1000", date(2019, 1, 1), date(2019, 1, 10), "handlowe")
    assert old.odcinki[0].stopa == Decimal("9.50")


def test_interest_merges_equal_rates_and_warns_after_check_date():
    w = odsetki.oblicz("1000", date(2024, 1, 1), date(2024, 12, 31), "kapitalowe")
    assert len(w.odcinki) == 1 and w.odcinki[0].dni == 366 and w.odcinki[0].stopa == Decimal("9.25")
    future = odsetki.oblicz("1000", date(2026, 10, 1), date(2027, 1, 31))
    assert any("sprawdzona" in o for o in future.ostrzezenia)


@pytest.mark.parametrize("args", [("0", date(2025, 1, 1), date(2025, 1, 2)), ("100", date(2015, 1, 1), date(2016, 2, 1)),
                                  ("100", date(2025, 2, 1), date(2025, 1, 1)), ("abc", date(2025, 1, 1), date(2025, 1, 2))])
def test_interest_rejects_bad_input(args):
    with pytest.raises(odsetki.OdsetkiError):
        odsetki.oblicz(*args)


def test_rate_table_sorted_and_parse_date():
    dates = [d for d, _ in odsetki.STOPA_REFERENCYJNA]
    assert dates == sorted(dates) and len(set(dates)) == len(dates)
    assert odsetki.parse_date("05.03.2026") == odsetki.parse_date("2026-03-05") == date(2026, 3, 5)
    with pytest.raises(odsetki.OdsetkiError):
        odsetki.parse_date("31.02.2026")


# ---------- bramka narzędzi i odpowiedź modelu ----------

@pytest.mark.parametrize("name,inp,web,ok", [
    ("mcp__kancelaria__oblicz_odsetki", {}, True, True),
    ("WebSearch", {"query": "art. 210 ksh", "allowed_domains": ["isap.sejm.gov.pl", "sn.pl"]}, True, True),
    ("WebSearch", {"query": "x"}, True, False),
    ("WebSearch", {"query": "x", "allowed_domains": ["isap.sejm.gov.pl", "evil.com"]}, True, False),
    ("WebSearch", {"query": "x", "allowed_domains": ["sn.pl"]}, False, False),
    ("WebFetch", {"url": "https://isap.sejm.gov.pl/isap.nsf/x"}, True, True),
    ("WebFetch", {"url": "https://www.sn.pl/orzecznictwo"}, True, True),
    ("WebFetch", {"url": "https://sn.pl.evil.com/x"}, True, False),
    ("WebFetch", {"url": "https://user@isap.sejm.gov.pl/"}, True, False),
    ("WebFetch", {"url": "file:///etc/passwd"}, True, False),
    ("Bash", {"command": "ls"}, True, False),
    ("mcp__figma__whoami", {}, True, False),
])
def test_tool_gate(name, inp, web, ok):
    assert decide(name, inp, web)[0] is ok


def test_parse_reply_takes_json_with_message():
    text = 'Myślę… {"a": 1} i wynik: {"wiadomosc": "Cześć Basiu", "status": "gotowe"}'
    assert parse_reply(text)["wiadomosc"] == "Cześć Basiu"
    assert parse_reply("bez jsona")["wymaga_czlowieka"] is True


def test_system_prompt_mentions_rules():
    sp = system_prompt(True, True)
    assert "Monika" in sp and "NIGDY nie podawaj sygnatury" in sp and "isap.sejm.gov.pl" in sp
    assert "wyłączone" in system_prompt(False, False)


def test_planner_tools_and_untrusted_report(make_env):
    pytest.importorskip("claude_agent_sdk")
    from kancelaria_agent.planner_agent import AgentPlanner
    from kancelaria_agent.projects import Projects
    cfg, store, *_ = make_env()
    projects = Projects(store, cfg.projects_dir)
    planner = AgentPlanner(BASIA, api_key="k", projects=projects, templates_dir=cfg.templates_dir)
    p = projects.create(BASIA, "x")
    names = {t.name for t in planner._tools(p["id"])}
    assert {"oblicz_odsetki", "czytaj_dokument", "utworz_pismo", "wzory_lista", "pamiec_szukaj"} <= names
    assert not any(n.startswith("wp_") or "figma" in n for n in names)
    opts = planner._options(p["id"])
    assert set(opts.tools) == {"WebSearch", "WebFetch"} and list(opts.mcp_servers) == ["kancelaria"]
    acts = planner.plan("email", {"trusted": False, "sender": "obcy@x.pl", "subject": "s", "body": "b",
                                  "message_id": "m"})
    assert [a.kind for a in acts] == [guards.SEND_MAIL] and acts[0].payload["to"] == BASIA


def test_planner_reply_goes_to_sender_and_filters_attachments(make_env, monkeypatch):
    pytest.importorskip("claude_agent_sdk")
    from kancelaria_agent.planner_agent import AgentPlanner
    from kancelaria_agent.projects import Projects
    cfg, store, *_ = make_env()
    projects = Projects(store, cfg.projects_dir)
    planner = AgentPlanner(BASIA, api_key="k", projects=projects)
    p = projects.create(ANETA, "x")
    (projects.dir(p) / "pisma" / "a.docx").write_bytes(b"x")
    monkeypatch.setattr(planner, "run_agent", lambda prompt, pid=None: json.dumps(
        {"wiadomosc": "Cześć Aneto", "zalaczniki": ["pisma/a.docx", "pisma/brak.docx", "../config.yaml"],
         "status": "gotowe"}))
    acts = planner.plan("email", {"trusted": True, "sender": ANETA, "subject": "s", "body": "b",
                                  "message_id": "m", "project_id": p["id"]})
    mail = [a for a in acts if a.kind == guards.SEND_MAIL][0].payload
    assert mail["to"] == ANETA and mail["project_attachments"] == ["pisma/a.docx"]
