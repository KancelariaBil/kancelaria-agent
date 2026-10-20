"""Dokumenty sprawy: odczyt (PDF, Word, Excel, tekst, obrazy) i tworzenie pism .docx.

Bezpieczeństwo: pliki to dane – nic z nich nie jest wykonywane (makra, skrypty, linki zewnętrzne).
Konwersje starszych formatów (.doc, .odt, .rtf) robi LibreOffice bez sieci, w katalogu tymczasowym,
z limitem czasu. Renderowanie stron PDF do obrazów (skany) – pdftoppm z limitem stron.
"""

from __future__ import annotations

import base64
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

MAX_TEXT = 120_000          # znaków tekstu zwracanych modelowi na raz
MAX_IMAGE = 5 * 1024 * 1024
SCAN_PAGES = 6              # ile stron skanu renderujemy na raz
TEXT_EXTS = {".txt", ".csv", ".xml", ".md", ".eml", ".json", ".html", ".htm"}
IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}


class DocumentError(ValueError):
    pass


# ---------- odczyt ----------

def _clip(text: str, offset: int = 0) -> str:
    text = text[offset:]
    if len(text) > MAX_TEXT:
        return text[:MAX_TEXT] + f"\n[… dalsza część od znaku {offset + MAX_TEXT} – wywołaj ponownie z od_znaku]"
    return text


def docx_text(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    out = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t"))
            out.append(text)
        elif tag == "tbl":
            for row in child.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iterchildren():
                        if cell.tag.endswith("}tc"):
                            cells.append(" ".join("".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                                                  for p in cell.iter() if p.tag.endswith("}p")).strip())
                    out.append(" | ".join(cells))
    return "\n".join(out).strip()


def pdf_text(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:   # noqa: BLE001 – uszkodzona strona nie przerywa odczytu reszty
            t = ""
        pages.append(f"--- strona {i} ---\n{t.strip()}")
    return "\n".join(pages), len(reader.pages)


def pdf_page_images(path: Path, first: int, count: int = SCAN_PAGES) -> list[bytes]:
    """Strony PDF jako JPEG (skany bez warstwy tekstu)."""
    if not shutil.which("pdftoppm"):
        raise DocumentError("brak pdftoppm na serwerze – skanu nie da się obejrzeć")
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["pdftoppm", "-jpeg", "-r", "110", "-f", str(first), "-l", str(first + count - 1),
                        str(path), f"{d}/s"], check=True, timeout=120, capture_output=True)
        return [p.read_bytes() for p in sorted(Path(d).glob("s*.jpg"))]


def convert_with_libreoffice(path: Path, target: str = "docx") -> Path:
    """Kopia .doc/.odt/.rtf jako .docx (obok, w katalogu tymczasowym – zwraca ścieżkę w tmp)."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        raise DocumentError("brak LibreOffice na serwerze – tego formatu nie odczytam")
    out = Path(tempfile.mkdtemp(prefix="konw-"))
    profile = Path(tempfile.mkdtemp(prefix="lo-"))
    try:
        subprocess.run([office, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                        "--convert-to", target, "--outdir", str(out), str(path)],
                       check=True, timeout=180, capture_output=True)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    files = list(out.glob(f"*.{target}"))
    if not files:
        raise DocumentError("konwersja nie powiodła się")
    return files[0]


def xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- arkusz {ws.title} ---")
        for row in ws.iter_rows(values_only=True, max_row=2000):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def read_document(path: Path, offset: int = 0, page: int | None = None) -> list[dict]:
    """Treść dokumentu jako bloki MCP: tekst albo obrazy (skan, zdjęcie)."""
    ext = path.suffix.lower()
    text = lambda t: [{"type": "text", "text": t}]   # noqa: E731
    if ext in IMAGE_TYPES:
        if path.stat().st_size > MAX_IMAGE:
            raise DocumentError("obraz większy niż 5 MB")
        return [{"type": "image", "data": base64.b64encode(path.read_bytes()).decode(), "mimeType": IMAGE_TYPES[ext]}]
    if ext == ".pdf":
        if page:
            imgs = pdf_page_images(path, page)
            return [{"type": "text", "text": f"strony od {page} (obrazy):"}] + [
                {"type": "image", "data": base64.b64encode(b).decode(), "mimeType": "image/jpeg"} for b in imgs]
        content, n = pdf_text(path)
        letters = len(re.sub(r"[\W\d_]+|strona", "", content))
        if letters < 40 * n:   # mało tekstu na stronę – najpewniej skan
            imgs = pdf_page_images(path, 1)
            note = (f"PDF ({n} str.) bez warstwy tekstu (skan) – poniżej strony 1–{len(imgs)} jako obrazy; "
                    f"kolejne: czytaj_dokument ze strona={len(imgs) + 1}.")
            return [{"type": "text", "text": note}] + [
                {"type": "image", "data": base64.b64encode(b).decode(), "mimeType": "image/jpeg"} for b in imgs]
        return text(f"PDF, {n} str.\n" + _clip(content, offset))
    if ext == ".docx":
        return text(_clip(docx_text(path), offset))
    if ext in (".doc", ".odt", ".rtf"):
        conv = convert_with_libreoffice(path)
        try:
            return text(_clip(docx_text(conv), offset))
        finally:
            shutil.rmtree(conv.parent, ignore_errors=True)
    if ext in (".xlsx", ".xlsm"):
        return text(_clip(xlsx_text(path), offset))
    if ext == ".xls":
        conv = convert_with_libreoffice(path, "xlsx")
        try:
            return text(_clip(xlsx_text(conv), offset))
        finally:
            shutil.rmtree(conv.parent, ignore_errors=True)
    if ext in TEXT_EXTS:
        return text(_clip(path.read_text(encoding="utf-8", errors="replace"), offset))
    raise DocumentError(f"nie odczytam pliku typu {ext or '(bez rozszerzenia)'}")


# ---------- tworzenie pism ----------

_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_runs(par, text: str) -> None:
    for chunk in _INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            par.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            par.add_run(chunk[1:-1]).italic = True
        else:
            par.add_run(chunk)


def _style(doc, *names):
    for n in names:
        try:
            return doc.styles[n]
        except KeyError:
            continue
    return None


def build_docx(content: str, target: Path, template: Path | None = None) -> Path:
    """Pismo z prostego zapisu: '# ' / '## ' nagłówki, '- ' punktory, '1. ' numeracja, '**pogrubienie**',
    '| a | b |' tabele, '---' podział strony, '>> ' akapit do prawej (miejscowość, data). Wzór (.docx) daje
    style, nagłówek i stopkę kancelarii – jego treść główna jest zastępowana nową."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

    if template is not None:
        doc = Document(str(template))
        body = doc.element.body
        for child in list(body.iterchildren()):
            if not child.tag.endswith("}sectPr"):   # układ strony, nagłówek i stopka zostają
                body.remove(child)
    else:
        doc = Document()
    lines = str(content).replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|") and line.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                st = _style(doc, "Table Grid", "Tabela - Siatka")
                if st is not None:
                    table.style = st
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        _add_runs(table.cell(r, c).paragraphs[0], val)
            continue
        if line.strip() == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-•]\s+", line):
            par = doc.add_paragraph(style=_style(doc, "List Bullet", "Lista punktowana"))
            _add_runs(par, re.sub(r"^\s*[-•]\s+", "", line))
        elif re.match(r"^\s*(\d+[.)]|§\s*\d+)", line):
            # Numeracja zostaje w tekście (§ 1, 1., 2)) – w pismach liczy się dokładnie to, co napisano.
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(par, line.strip())
        elif line.startswith(">> "):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _add_runs(par, line[3:])
        elif not line.strip():
            pass
        else:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_runs(par, line)
        i += 1
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target
